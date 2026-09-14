from __future__ import annotations

import csv
import subprocess
from io import StringIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT_PATH = Path("exports/thinkspace_reminder_candidates_20260910.xlsx")


QUERY = r"""
COPY (
WITH learner_flags AS (
    SELECT
        u.registration_id,
        u.full_name,
        u.email,
        u.role,
        COALESCE(NULLIF(u.team_name, ''), 'Cá nhân tự do') AS project_name,
        u.has_ever_accessed,
        u.last_access_at,
        COALESCE(u.viewed_activity_count, 0) AS viewed_activity_count,
        COALESCE(u.learning_event_count, 0) AS learning_event_count,
        COUNT(*) FILTER (
            WHERE e.moodle_course_module_id = 707
              AND e.is_access_event = TRUE
        ) AS pre_program_survey_view_count,
        COUNT(DISTINCT e.moodle_course_module_id) FILTER (
            WHERE e.moodle_course_module_id IN (
                650, 653, 656, 659, 662, 665,
                651, 654, 657, 660, 663, 666,
                714, 716
            )
              AND e.is_access_event = TRUE
        ) AS milestone_fmc3_module_count,
        COUNT(*) FILTER (
            WHERE e.moodle_course_module_id IN (
                650, 653, 656, 659, 662, 665,
                651, 654, 657, 660, 663, 666,
                714, 716
            )
              AND e.is_access_event = TRUE
        ) AS milestone_fmc3_view_count
    FROM gold_registered_user_learning_summary u
    LEFT JOIN silver_moodle_learning_events e
        ON LOWER(TRIM(e.email)) = u.email
    GROUP BY
        u.registration_id,
        u.full_name,
        u.email,
        u.role,
        COALESCE(NULLIF(u.team_name, ''), 'Cá nhân tự do'),
        u.has_ever_accessed,
        u.last_access_at,
        u.viewed_activity_count,
        u.learning_event_count
)
SELECT
    registration_id,
    full_name,
    email,
    role,
    project_name,
    has_ever_accessed,
    viewed_activity_count,
    learning_event_count,
    pre_program_survey_view_count,
    milestone_fmc3_module_count,
    milestone_fmc3_view_count,
    last_access_at
FROM learner_flags
WHERE has_ever_accessed = FALSE
   OR pre_program_survey_view_count = 0
   OR milestone_fmc3_module_count = 0
ORDER BY
    CASE
        WHEN has_ever_accessed = FALSE THEN 1
        WHEN pre_program_survey_view_count = 0
         AND milestone_fmc3_module_count = 0 THEN 1
        WHEN pre_program_survey_view_count = 0 THEN 2
        WHEN milestone_fmc3_module_count = 0 THEN 2
        ELSE 3
    END,
    project_name,
    full_name,
    email
) TO STDOUT WITH CSV HEADER
"""


HEADERS = [
    ("registration_id", "Registration ID"),
    ("full_name", "Họ và tên"),
    ("email", "Email"),
    ("role", "Vai trò"),
    ("project_name", "Dự án"),
    ("thinkspace_access_status", "Trạng thái ThinkSpace"),
    ("viewed_activity_count", "Số hoạt động đã xem"),
    ("learning_event_count", "Log events"),
    ("pre_program_survey_view_count", "Lượt vào Pre-Program Survey"),
    ("milestone_fmc3_module_count", "Số module milestone/FMC3 đã vào"),
    ("milestone_fmc3_view_count", "Lượt xem milestone/FMC3"),
    ("last_access_at", "Lần hoạt động cuối"),
    ("reminder_priority", "Mức ưu tiên"),
    ("reminder_reason", "Lý do cần remind"),
]


def run_query() -> list[dict[str, str]]:
    result = subprocess.run(
        [
            "docker",
            "exec",
            "-e",
            "PGCLIENTENCODING=UTF8",
            "thinkspace_management-db-1",
            "psql",
            "-U",
            "thinkspace",
            "-d",
            "thinkspacedb",
            "-c",
            QUERY,
        ],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    return list(csv.DictReader(StringIO(result.stdout)))


def decorate_row(row: dict[str, str]) -> dict[str, str]:
    has_accessed = row.get("has_ever_accessed") == "t"
    survey_views = int(row.get("pre_program_survey_view_count") or 0)
    milestone_modules = int(row.get("milestone_fmc3_module_count") or 0)

    reasons: list[str] = []
    if not has_accessed:
        reasons.append("Chưa có bất kỳ hoạt động nào trên ThinkSpace/Moodle")
    if has_accessed and survey_views == 0:
        reasons.append("Đã vào ThinkSpace nhưng chưa vào Pre-Program Survey")
    elif not has_accessed:
        reasons.append("Chưa vào Pre-Program Survey")
    if milestone_modules == 0:
        reasons.append("Chưa có hoạt động ở các milestone hoặc khóa entrepreneurship/FMC3")

    if not has_accessed:
        priority = "Cao"
    elif survey_views == 0 and milestone_modules == 0:
        priority = "Cao"
    elif survey_views == 0 or milestone_modules == 0:
        priority = "Vừa"
    else:
        priority = "Thấp"

    return {
        **row,
        "thinkspace_access_status": "Đã có hoạt động" if has_accessed else "Chưa có hoạt động",
        "reminder_priority": priority,
        "reminder_reason": "; ".join(reasons),
    }


def write_workbook(rows: list[dict[str, str]]) -> None:
    OUTPUT_PATH.parent.mkdir(exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Reminder candidates"
    ws.append([label for _, label in HEADERS])

    for raw_row in rows:
        row = decorate_row(raw_row)
        ws.append([row.get(key, "") for key, _ in HEADERS])

    header_fill = PatternFill("solid", fgColor="254385")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)

    for worksheet_row in ws.iter_rows(min_row=2):
        priority = worksheet_row[12].value
        fill = None
        if priority == "Cao":
            fill = PatternFill("solid", fgColor="FCE4D6")
        elif priority == "Vừa":
            fill = PatternFill("solid", fgColor="FFF2CC")

        for cell in worksheet_row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if fill:
                cell.fill = fill

    widths = [14, 28, 34, 16, 42, 24, 18, 14, 24, 26, 22, 24, 14, 68]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    decorated_rows = [decorate_row(row) for row in rows]
    summary_rows = [
        ("Tổng dòng cần remind", len(decorated_rows)),
        ("Ưu tiên Cao", sum(1 for row in decorated_rows if row["reminder_priority"] == "Cao")),
        ("Ưu tiên Vừa", sum(1 for row in decorated_rows if row["reminder_priority"] == "Vừa")),
        (
            "Chưa có hoạt động ThinkSpace",
            sum(1 for row in decorated_rows if row["thinkspace_access_status"] == "Chưa có hoạt động"),
        ),
        (
            "Đã vào nhưng chưa vào Pre-Program Survey",
            sum(
                1
                for row in decorated_rows
                if row["thinkspace_access_status"] == "Đã có hoạt động"
                and int(row["pre_program_survey_view_count"] or 0) == 0
            ),
        ),
        (
            "Chưa có hoạt động milestone/FMC3",
            sum(1 for row in decorated_rows if int(row["milestone_fmc3_module_count"] or 0) == 0),
        ),
    ]

    summary = wb.create_sheet("Summary")
    summary.append(["Chỉ số", "Giá trị"])
    for item in summary_rows:
        summary.append(list(item))
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    summary.column_dimensions["A"].width = 48
    summary.column_dimensions["B"].width = 18

    wb.save(OUTPUT_PATH)


def main() -> None:
    rows = run_query()
    write_workbook(rows)
    print(OUTPUT_PATH.resolve())
    print(f"Tong dong can remind: {len(rows)}")


if __name__ == "__main__":
    main()
